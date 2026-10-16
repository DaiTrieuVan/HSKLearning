from pathlib import Path
from collections import defaultdict
import argparse, re
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def rows(sheet):
    it = sheet.iter_rows(values_only=True)
    headers = next(it)
    return [dict(zip(headers, r)) for r in it if any(v is not None for v in r)]

def val(row, index):
    return list(row.values())[index] if index < len(row) else None

def clean(x):
    return str(x).strip() if x is not None else ''

def section(title, fields):
    return '\n## ' + title + '\n\n' + ''.join(f'**{name}:** {clean(value)}\n\n' for name,value in fields if clean(value))

def main():
    p = argparse.ArgumentParser()
    p.add_argument('--input', default=str(ROOT/'docs'/'requirements.xlsx'))
    p.add_argument('--output', default=str(ROOT/'docs'/'user-stories'))
    a=p.parse_args()
    wb=load_workbook(a.input, read_only=True, data_only=True)
    stories=rows(wb['Product Backlog (User Stories)'])
    ac=defaultdict(list); tasks=defaultdict(list); rules={}
    for r in rows(wb['Acceptance Criteria']): ac[clean(val(r,1))].append(r)
    for r in rows(wb['Tasks (Công việc)']): tasks[clean(val(r,1))].append(r)
    for r in rows(wb['Business Rules (Quy tắc)']): rules[clean(val(r,0))]=r
    out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
    generated=set(); warnings=[]
    for s in stories:
        sid=clean(val(s,3))
        if not re.fullmatch(r'HTT-\d{2}-CN-\d{3}',sid):
            warnings.append(f'Mã story không hợp lệ: {sid!r}');continue
        if sid in generated:
            warnings.append(f'Trùng story: {sid}');continue
        generated.add(sid)
        lines=[f'# {sid} — {clean(val(s,4))}\n',f'> Epic: {clean(val(s,1))} — {clean(val(s,2))}\n']
        lines.append(section('Thông tin User Story',[(k,val(s,i)) for k,i in [('User Story',5),('Vai trò',6),('Giá trị nghiệp vụ',7),('Mô tả chi tiết',8),('Điều kiện bắt đầu',9),('Kết quả hoàn thành',10),('Quy tắc liên quan',11),('Độ ưu tiên',12),('Story Point',13),('Phụ thuộc',14),('Trạng thái',15),('Ghi chú',16)]]))
        rule_ids=re.findall(r'QTN-\d+',clean(val(s,11)))
        rule_ids += [rid for rid,r in rules.items() if sid in clean(val(r,8))]
        lines.append('\n## Quy tắc nghiệp vụ liên quan\n')
        for rid in dict.fromkeys(rule_ids):
            r=rules.get(rid)
            if not r: warnings.append(f'{sid}: không tìm thấy {rid}');continue
            lines.append(f'\n### {rid} — {clean(val(r,1))}\n')
            lines.append(section('Chi tiết quy tắc',[(k,val(r,i)) for k,i in [('Mô tả',2),('Vai trò',3),('Đối tượng',4),('Điều kiện',5),('Khi thỏa mãn',6),('Khi không thỏa mãn',7)]]))
        lines.append('\n## Acceptance Criteria / Test Cases\n')
        for r in ac[sid]:
            lines.append(f'\n### {clean(val(r,2))} — {clean(val(r,3))}\n')
            lines.append(section('Given / When / Then',[(k,val(r,i)) for k,i in [('Given',4),('When',5),('Then',6),('Test Data',7),('Severity',8)]]))
        lines.append('\n## Công việc (Tasks)\n')
        if not tasks[sid]:
            lines.append('\n_Chưa có Task được định nghĩa trong tài liệu nguồn. Agent có thể đề xuất Implementation Tasks riêng, nhưng không được coi đó là Task đã được phê duyệt._\n')
        for r in tasks[sid]:
            lines.append(f'\n### {clean(val(r,2))} — {clean(val(r,3))}\n')
            lines.append(section('Chi tiết công việc',[(k,val(r,i)) for k,i in [('Kết quả cần đạt',4),('Nhóm công việc',5),('Độ ưu tiên',6),('Story Point',7),('Trạng thái',8),('Ghi chú',9)]]))
        if not ac[sid]: warnings.append(f'{sid}: không có test case')
        (out/f'{sid}.md').write_text('\n'.join(lines),encoding='utf-8')
    for sid in set(ac)|set(tasks):
        if sid not in generated: warnings.append(f'Dữ liệu tham chiếu story không tồn tại: {sid}')
    print(f'Đã xuất {len(generated)} User Stories, {sum(map(len,ac.values()))} Test Cases, {sum(map(len,tasks.values()))} Tasks')
    if warnings:
        print('CẢNH BÁO:')
        for w in warnings: print('-',w)

if __name__=='__main__':main()
